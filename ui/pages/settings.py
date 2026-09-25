import tkinter as tk
from tkinter import ttk, messagebox
import logging
from config import (
    ACCENT_COLOR, ACCENT_LIGHT, FONT_FAMILY, APP_NAME, VERSION,
    USER_DATA_DIR, update_data_dir, TEXT_PRIMARY, TEXT_SECONDARY, TEXT_MUTED,
    FONT_SIZE_MD, FONT_SIZE_LG, FONT_SIZE_XL, FONT_SIZE_XXL,
    get_setting, set_setting, CARD_BG, SUCCESS_COLOR, DANGER_COLOR, WARNING_COLOR,
)
from utils.company import load_company, save_company
from utils.update_checker import get_update_status
from config import RELEASE_BASE_URL
from database.backup import start_auto_backup, stop_auto_backup
from ui.pages.license_page import build_license_tab  # noqa: F401


_log = logging.getLogger(__name__)


class SettingsPage(ttk.Frame):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        self._app = self.winfo_toplevel()
        self._build_ui()

    def _build_ui(self):
        header = ttk.Label(self, text="Settings",
                           font=(FONT_FAMILY, 20, "bold"))
        header.pack(anchor="w", padx=20, pady=(20, 10))

        self._notebook = ttk.Notebook(self)
        self._notebook.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 20))

        tabs = [
            ("  General  ", self._general_tab),
            ("  Business  ", self._business_tab),
            ("  Vertical  ", self._vertical_tab),
            ("  Features  ", self._features_tab),
            ("  Tax  ", self._tax_tab),
            ("  Notifications  ", self._notifications_tab),
            ("  Users  ", self._users_tab),
            ("  Audit Log  ", self._audit_log_tab),
            ("  Company  ", self._company_tab),
            ("  Backup  ", self._backup_tab),
            ("  License  ", self._license_tab),
            ("  Updates  ", self._updates_tab),
        ]
        for label, builder in tabs:
            try:
                builder(self._notebook)
            except Exception as exc:
                _log.exception("Failed to build settings tab %s", label)
                self._add_error_tab(label.strip(), exc)

    def _add_error_tab(self, label, exc):
        frame = ttk.Frame(self._notebook, padding=20)
        self._notebook.add(frame, text=label)
        tk.Label(frame, text=f"⚠  This tab failed to load.",
                 font=(FONT_FAMILY, 12, "bold"),
                 bg=CARD_BG, fg="#DC2626").pack(anchor="w")
        tk.Label(frame, text=str(exc),
                 font=(FONT_FAMILY, 9), bg=CARD_BG,
                 fg=TEXT_MUTED, wraplength=600, justify="left").pack(anchor="w", pady=(4, 0))

    def _general_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  General  ")

        ttk.Label(frame, text="Currency Symbol",
                  font=(FONT_FAMILY, 11)).grid(row=0, column=0, sticky="w", pady=8)
        self.currency_var = tk.StringVar(value=get_setting("currency_symbol", "\u20B9"))
        ttk.Entry(frame, textvariable=self.currency_var, width=10).grid(
            row=0, column=1, sticky="w", padx=10)

        ttk.Label(frame, text="Theme",
                  font=(FONT_FAMILY, 11)).grid(row=1, column=0, sticky="w", pady=8)
        self.theme_var = tk.StringVar(value=get_setting("theme", "Light"))
        ttk.Combobox(frame, textvariable=self.theme_var,
                     values=["Light", "Dark"], state="readonly", width=12).grid(
            row=1, column=1, sticky="w", padx=10)
        self.theme_var.trace_add("write", self._toggle_theme)

        ttk.Separator(frame, orient="horizontal").grid(
            row=2, column=0, columnspan=3, sticky="ew", pady=10)

        ttk.Label(frame, text="Data Location",
                  font=(FONT_FAMILY, 11)).grid(row=3, column=0, sticky="nw", pady=8)
        self.data_dir_var = tk.StringVar(value=str(USER_DATA_DIR))
        data_entry = ttk.Entry(frame, textvariable=self.data_dir_var, width=45)
        data_entry.grid(row=3, column=1, sticky="w", padx=10)
        ttk.Button(frame, text="Browse...",
                   command=self._browse_data_dir).grid(row=3, column=2, padx=5)

        ttk.Separator(frame, orient="horizontal").grid(
            row=4, column=0, columnspan=3, sticky="ew", pady=10)

        # ---- Per-user preferences ----
        from utils.settings_helper import get_user_setting, set_user_setting
        try:
            from utils.auth import auth_manager
            _uname = auth_manager.get_current_user() or "current user"
        except Exception:
            _uname = "current user"

        ttk.Label(frame, text=f"My Preferences (per user: {_uname})",
                  font=(FONT_FAMILY, 11, "bold")).grid(
            row=5, column=0, columnspan=3, sticky="w", pady=(0, 6))

        ttk.Label(frame, text="Default Start Page",
                  font=(FONT_FAMILY, 10)).grid(row=6, column=0, sticky="w", pady=4)
        from ui.app import NAV_ITEMS
        nav_keys = [k for k, _ in NAV_ITEMS]
        nav_labels = {k: n.split("  ")[-1] for k, n in NAV_ITEMS}
        cur_start = get_user_setting("nav.start_page", "dashboard")
        if cur_start not in nav_keys:
            cur_start = "dashboard"
        self._start_page_var = tk.StringVar(value=nav_labels.get(cur_start, "Dashboard"))
        start_combo = ttk.Combobox(frame, textvariable=self._start_page_var,
                                   values=[nav_labels[k] for k in nav_keys],
                                   state="readonly", width=18)
        start_combo.grid(row=6, column=1, sticky="w", padx=10)
        self._start_page_lookup = {nav_labels[k]: k for k in nav_keys}

        def _save_start_page(*_):
            set_user_setting("nav.start_page", self._start_page_lookup.get(
                self._start_page_var.get(), "dashboard"))
            try:
                self._app.toast.show("Start page saved", "success", 2500)
            except Exception:
                pass
        start_combo.bind("<<ComboboxSelected>>", _save_start_page)

        ttk.Label(frame, text="Dashboard Cards",
                  font=(FONT_FAMILY, 10)).grid(row=7, column=0, sticky="nw", pady=4)
        cards_frame = ttk.Frame(frame)
        cards_frame.grid(row=7, column=1, sticky="w", padx=10)
        from utils.settings_helper import is_feature_enabled as _ife
        all_cards = [
            ("total_items", "Total Items"), ("low_stock", "Low Stock"),
            ("stock_value", "Stock Value"), ("sales_today", "Today's Revenue"),
            ("monthly_revenue", "Monthly Revenue"), ("total_pending", "Pending Payments"),
            ("extra_income_today", "Extra Income Today"), ("total_customers", "Total Customers"),
        ]
        if _ife("expiry_tracking"):
            all_cards.append(("expiring_soon", "Expiring Soon"))
        hidden = set(get_user_setting("dashboard.hidden_cards", []) or [])
        self._card_vars = {}
        for i, (key, label) in enumerate(all_cards):
            var = tk.BooleanVar(value=key not in hidden)
            var._key = key
            self._card_vars[key] = var
            tk.Checkbutton(cards_frame, text=label, variable=var,
                           command=self._save_card_prefs).grid(
                row=i // 2, column=i % 2, sticky="w", padx=(0, 14))

        ttk.Label(frame, text="Card visibility applies to your login only.",
                  font=(FONT_FAMILY, 8), foreground=TEXT_MUTED).grid(
            row=8, column=1, sticky="w", padx=10)

        # ---- Dashboard card arrangement (drag-free: Up/Down ordering) ----
        ttk.Label(frame, text="Card Order",
                  font=(FONT_FAMILY, 10)).grid(row=9, column=0, sticky="nw", pady=4)
        order_box = ttk.Frame(frame)
        order_box.grid(row=9, column=1, sticky="w", padx=10)
        self._card_order_lb = tk.Listbox(order_box, height=min(len(all_cards), 9), width=28,
                                         exportselection=False, font=(FONT_FAMILY, 9))
        self._card_order_lb.pack(side=tk.LEFT)
        ord_btns = ttk.Frame(order_box)
        ord_btns.pack(side=tk.LEFT, padx=(6, 0))
        ttk.Button(ord_btns, text="↑ Up", width=7, style="Secondary.TButton",
                   command=lambda: self._move_card(-1)).pack(pady=1)
        ttk.Button(ord_btns, text="↓ Down", width=7, style="Secondary.TButton",
                   command=lambda: self._move_card(1)).pack(pady=1)

        saved_order = get_user_setting("dashboard.card_order", []) or []
        label_for = dict(all_cards)
        ordered_keys = [k for k in saved_order if k in label_for]
        ordered_keys += [k for k, _ in all_cards if k not in ordered_keys]
        self._card_order_keys = ordered_keys
        for k in ordered_keys:
            self._card_order_lb.insert(tk.END, label_for[k])

        # ---- Sidebar page visibility ----
        ttk.Label(frame, text="Sidebar Pages",
                  font=(FONT_FAMILY, 10)).grid(row=10, column=0, sticky="nw", pady=4)
        nav_frame = ttk.Frame(frame)
        nav_frame.grid(row=10, column=1, sticky="w", padx=10)
        hidden_pages = set(get_user_setting("nav.hidden_pages", []) or [])
        self._nav_page_vars = {}
        for i, (key, label) in enumerate(NAV_ITEMS):
            if key == "settings":
                continue  # settings must always remain reachable
            var = tk.BooleanVar(value=key not in hidden_pages)
            var._key = key
            self._nav_page_vars[key] = var
            tk.Checkbutton(nav_frame, text=label.split("  ")[-1], variable=var,
                           command=self._save_nav_prefs).grid(
                row=i // 3, column=i % 3, sticky="w", padx=(0, 14))
        ttk.Label(frame, text="Bank Recon also needs its feature enabled (Features tab).",
                  font=(FONT_FAMILY, 8), foreground=TEXT_MUTED).grid(
            row=11, column=1, sticky="w", padx=10)

        ttk.Separator(frame, orient="horizontal").grid(
            row=12, column=0, columnspan=3, sticky="ew", pady=10)

        ttk.Label(frame, text="About",
                  font=(FONT_FAMILY, 11)).grid(row=13, column=0, sticky="w", pady=8)
        ttk.Label(frame,
                  text=f"{APP_NAME} v{VERSION}\n"
                       "Built with Python & Tkinter\n"
                       "Database: Microsoft Excel (.xlsx)",
                  foreground="#777").grid(row=13, column=1, sticky="w", padx=10)

    def _move_card(self, direction):
        lb = self._card_order_lb
        sel = lb.curselection()
        if not sel:
            return
        i = sel[0]
        j = i + direction
        if j < 0 or j >= lb.size():
            return
        key = self._card_order_keys[i]
        self._card_order_keys[i], self._card_order_keys[j] = \
            self._card_order_keys[j], self._card_order_keys[i]
        txt = lb.get(i)
        lb.delete(i)
        lb.insert(j, txt)
        lb.selection_set(j)
        self._save_card_prefs()

    def _save_nav_prefs(self):
        from utils.settings_helper import set_user_setting
        hidden = [k for k, var in self._nav_page_vars.items() if not var.get()]
        set_user_setting("nav.hidden_pages", hidden)
        try:
            self._app._apply_sidebar_visibility()
            self._app.toast.show("Sidebar updated", "success", 2500)
        except Exception:
            pass

    def _save_card_prefs(self):
        from utils.settings_helper import set_user_setting
        hidden = [k for k, var in self._card_vars.items() if not var.get()]
        set_user_setting("dashboard.hidden_cards", hidden)
        if hasattr(self, "_card_order_keys"):
            set_user_setting("dashboard.card_order", list(self._card_order_keys))
        try:
            self._app.toast.show("Dashboard layout saved — open Dashboard to see it",
                                 "success", 3000)
        except Exception:
            pass

    def _backup_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Backup  ")

        ttk.Label(frame, text="Automatic Backup",
                  font=(FONT_FAMILY, 12, "bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))

        ttk.Label(frame, text="Backup Interval (minutes):",
                  font=(FONT_FAMILY, 10)).grid(
            row=1, column=0, sticky="w", pady=5, padx=(0, 10))
        self.backup_interval_var = tk.IntVar(value=get_setting("backup_interval_minutes", 30))
        interval_combo = ttk.Combobox(frame, textvariable=self.backup_interval_var,
                                       values=[15, 30, 60, 120, 180],
                                       state="readonly", width=10)
        interval_combo.grid(row=1, column=1, sticky="w", pady=5)
        interval_combo.bind("<<ComboboxSelected>>", self._on_backup_interval_change)

        ttk.Label(frame, text="How often to auto-backup data files.\n"
                              "Recommended: 30 minutes.",
                  font=(FONT_FAMILY, 9), foreground=TEXT_MUTED).grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(0, 10))

        ttk.Separator(frame, orient="horizontal").grid(
            row=3, column=0, columnspan=2, sticky="ew", pady=10)

        ttk.Label(frame, text="Manual Backup",
                  font=(FONT_FAMILY, 12, "bold")).grid(
            row=4, column=0, columnspan=2, sticky="w", pady=(0, 10))

        def do_manual_backup():
            from database.backup import create_backup
            if create_backup():
                messagebox.showinfo("Backup", "Backup created successfully")
            else:
                messagebox.showerror("Backup", "Backup failed")

        ttk.Button(frame, text="\u21BA  Create Backup Now",
                   command=do_manual_backup).grid(
            row=5, column=0, columnspan=2, sticky="w", pady=5)

        ttk.Label(frame, text="Creates a timestamped backup of all data files.",
                  font=(FONT_FAMILY, 9), foreground=TEXT_MUTED).grid(
            row=6, column=0, columnspan=2, sticky="w", pady=(0, 5))

        ttk.Separator(frame, orient="horizontal").grid(
            row=7, column=0, columnspan=2, sticky="ew", pady=12)

        # ---- Cloud backup ----
        from utils import cloud_backup
        from utils.license import license_mgr

        ttk.Label(frame, text="Cloud Backup",
                  font=(FONT_FAMILY, 12, "bold")).grid(
            row=8, column=0, columnspan=2, sticky="w", pady=(0, 6))

        if not license_mgr.has_feature("cloud_backup"):
            ttk.Label(frame, text="⚠  Cloud backup requires a Basic/Pro/Enterprise license "
                                  "(Settings → License).",
                      font=(FONT_FAMILY, 9), foreground=WARNING_COLOR,
                      wraplength=560).grid(row=9, column=0, columnspan=2, sticky="w", pady=(0, 6))

        self._cloud_enabled_var = tk.BooleanVar(value=get_setting("cloud_backup.enabled", False))
        tk.Checkbutton(frame, text="Enable cloud backup", variable=self._cloud_enabled_var,
                       font=(FONT_FAMILY, 10)).grid(row=10, column=0, columnspan=2, sticky="w")

        ttk.Label(frame, text="Endpoint URL (blank = license server default):",
                  font=(FONT_FAMILY, 10)).grid(row=11, column=0, sticky="w", pady=(8, 2))
        self._cloud_endpoint_var = tk.StringVar(value=get_setting("cloud_backup.endpoint", "") or "")
        ttk.Entry(frame, textvariable=self._cloud_endpoint_var, width=46).grid(
            row=11, column=1, sticky="w", pady=(8, 2))

        ttk.Label(frame, text="API Token (optional, stored encrypted):",
                  font=(FONT_FAMILY, 10)).grid(row=12, column=0, sticky="w", pady=2)
        self._cloud_token_var = tk.StringVar(value=cloud_backup.load_token())
        ttk.Entry(frame, textvariable=self._cloud_token_var, width=46, show="●").grid(
            row=12, column=1, sticky="w", pady=2)

        ttk.Label(frame, text="Upload every (hours):",
                  font=(FONT_FAMILY, 10)).grid(row=13, column=0, sticky="w", pady=2)
        self._cloud_hours_var = tk.StringVar(value=str(get_setting("cloud_backup.interval_hours", 24)))
        ttk.Entry(frame, textvariable=self._cloud_hours_var, width=8).grid(
            row=13, column=1, sticky="w", pady=2)

        cloud_row = ttk.Frame(frame)
        cloud_row.grid(row=14, column=0, columnspan=2, sticky="w", pady=(10, 0))
        ttk.Button(cloud_row, text="Save Cloud Settings",
                   command=self._save_cloud_backup).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(cloud_row, text="☁  Backup to Cloud Now",
                   command=self._cloud_backup_now).pack(side=tk.LEFT)

        last = get_setting("cloud_backup.last_upload", "") or "never"
        ttk.Label(frame, text=f"Last cloud backup: {last}",
                  font=(FONT_FAMILY, 9), foreground=TEXT_MUTED).grid(
            row=15, column=0, columnspan=2, sticky="w", pady=(6, 0))

    def _save_cloud_backup(self):
        from utils import cloud_backup
        set_setting("cloud_backup.enabled", bool(self._cloud_enabled_var.get()))
        set_setting("cloud_backup.endpoint", self._cloud_endpoint_var.get().strip())
        try:
            set_setting("cloud_backup.interval_hours", int(self._cloud_hours_var.get().strip() or 24))
        except ValueError:
            pass
        cloud_backup.save_token(self._cloud_token_var.get().strip())
        messagebox.showinfo("Saved", "Cloud backup settings saved")

    def _cloud_backup_now(self):
        from utils import notifier  # noqa: F401  (keeps import pattern same)
        from utils import cloud_backup
        from utils.feature_gate import require_feature
        if not require_feature("cloud_backup", parent=self, friendly_name="Cloud Backup"):
            return
        self._save_cloud_backup()
        app = self._app
        try:
            app.toast.show("Uploading backup…", "info", 3000)
        except Exception:
            pass

        def _done(ok, msg):
            self.after(0, lambda: app.toast.show(msg, "success" if ok else "error", 5000))

        cloud_backup.upload_backup_async(callback=_done)

    def _on_backup_interval_change(self, event=None):
        interval = self.backup_interval_var.get()
        set_setting("backup_interval_minutes", interval)
        stop_auto_backup()
        start_auto_backup(interval)
        self._app.toast.show(f"Backup interval set to {interval} minutes", "success", 3000)

    def _license_tab(self, notebook):
        build_license_tab(notebook, self._app)

    def _updates_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Updates  ")

        row = 0
        ttk.Label(frame, text="Update Settings",
                  font=(FONT_FAMILY, 12, "bold")).grid(
            row=row, column=0, columnspan=2, sticky="w", pady=(0, 10))
        row += 1

        status = get_update_status()

        info_items = [
            ("Current Version", f"v{VERSION}"),
            ("Latest Available", f"v{status['latest_version']}" if status['latest_version'] else "Not checked"),
            ("Last Checked", status['last_check'] if status['last_check'] else "Never"),
            ("Online Status", "\u2713 Online" if status['is_online_now'] else "\u2717 Offline"),
        ]
        for label, value in info_items:
            ttk.Label(frame, text=label + ":",
                      font=(FONT_FAMILY, 10, "bold")).grid(
                row=row, column=0, sticky="w", pady=3, padx=(0, 15))
            ttk.Label(frame, text=value,
                      font=(FONT_FAMILY, 10)).grid(
                row=row, column=1, sticky="w", pady=3)
            row += 1

        row += 1
        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=row, column=0, columnspan=2, pady=10, sticky="w")

        def check_now():
            self._app.check_updates_now(True)

        ttk.Button(btn_frame, text="\u21BB  Check for Updates Now",
                   command=check_now).pack(side=tk.LEFT, padx=5)

        def open_releases():
            import webbrowser
            webbrowser.open(RELEASE_BASE_URL)

        ttk.Button(btn_frame, text="\u2197  View All Releases",
                   command=open_releases).pack(side=tk.LEFT, padx=5)

        history = status.get("update_history", [])
        if history and isinstance(history, list) and len(history) > 0:
            row += 2
            sep = ttk.Separator(frame, orient="horizontal")
            sep.grid(row=row, column=0, columnspan=3, sticky="ew", pady=10)
            row += 1

            ttk.Label(frame, text="Update History",
                      font=(FONT_FAMILY, 11, "bold")).grid(
                row=row, column=0, columnspan=2, sticky="w", pady=(0, 5))
            row += 1

            tree_frame = ttk.Frame(frame)
            tree_frame.grid(row=row, column=0, columnspan=2, sticky="nsew")
            frame.grid_rowconfigure(row, weight=1)
            frame.grid_columnconfigure(1, weight=1)

            cols = ["Version", "Detected At", "Highlights"]
            tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=5)
            for col in cols:
                tree.heading(col, text=col)
                tree.column(col, width=120)
            tree.column("Highlights", width=250)

            for h in reversed(history[-10:]):
                tree.insert("", tk.END, values=[
                    h.get("version", ""),
                    h.get("detected_at", ""),
                    h.get("changelog", "")[:60],
                ])

            scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=scroll.set)
            tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
            scroll.pack(side=tk.RIGHT, fill=tk.Y)

    def _company_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Company  ")
        company = load_company()

        fields = [
            ("Business Name", "name", 35),
            ("Address", "address", 45),
            ("City", "city", 25),
            ("State", "state", 25),
            ("Pincode", "pincode", 15),
            ("Phone", "phone", 20),
            ("Email", "email", 30),
            ("GSTIN", "gstin", 25),
            ("PAN", "pan", 20),
            ("UPI ID (India) — optional", "upi_id", 30),
            ("Payment instructions (eSewa/bank/wallet note)", "payment_instructions", 45),
            ("Invoice Prefix", "invoice_prefix", 10),
        ]

        self._company_fields = {}
        for i, (label, key, width) in enumerate(fields):
            ttk.Label(frame, text=label,
                      font=(FONT_FAMILY, 10)).grid(
                row=i, column=0, sticky="w", pady=4, padx=(0, 10))
            entry = ttk.Entry(frame, width=width)
            entry.insert(0, company.get(key, ""))
            entry.grid(row=i, column=1, sticky="w", pady=4)
            self._company_fields[key] = entry

        ttk.Label(frame, text="Invoice Footer Note",
                  font=(FONT_FAMILY, 10)).grid(
            row=len(fields), column=0, sticky="w", pady=8, padx=(0, 10))
        self.inv_note_text = tk.Text(frame, width=40, height=3,
                                     font=(FONT_FAMILY, 10))
        self.inv_note_text.insert("1.0", company.get("invoice_note", ""))
        self.inv_note_text.grid(row=len(fields), column=1, sticky="w", pady=8)

        ttk.Button(frame, text="\u2713  Save Company Details",
                   command=self._save_company).grid(
            row=len(fields) + 1, column=0, columnspan=2, pady=15)

    def _save_company(self):
        data = {}
        for key, entry in self._company_fields.items():
            data[key] = entry.get().strip()
        data["invoice_note"] = self.inv_note_text.get("1.0", "end-1c").strip()
        save_company(data)
        messagebox.showinfo("Saved", "Company details saved successfully")

    def _browse_data_dir(self):
        from tkinter import filedialog
        path = filedialog.askdirectory(title="Select Data Folder")
        if path:
            ok, err = update_data_dir(path)
            if ok:
                self.data_dir_var.set(path)
                messagebox.showinfo("Data Location",
                                    f"Data will now be saved to:\n{path}\n\n"
                                    "Existing files must be moved manually.")
            else:
                messagebox.showerror("Invalid Data Folder",
                                     err or "Could not use that location.")

    def _toggle_theme(self, *args):
        theme = self.theme_var.get()
        set_setting("theme", theme)
        try:
            from utils.settings_helper import set_user_setting
            set_user_setting("theme", theme)
        except Exception:
            pass
        app = self._app
        app.reload_current_page()
        app._apply_theme()
        try:
            app.toast.show(f"Theme changed to {theme}", "success", 3000)
        except tk.TclError:
            pass

    # ========== NEW TABS (Business, Vertical, Features, Tax) ==========

    def _business_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Business  ")
        tk.Label(frame, text="Invoice Settings", font=(FONT_FAMILY, 12, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        tk.Label(frame, text="Invoice No. Prefix", font=(FONT_FAMILY, 10)).pack(anchor="w", pady=(8, 2))
        self._inv_pref_var = tk.StringVar(value=get_setting("business.invoice_no_prefix", "INV"))
        tk.Entry(frame, textvariable=self._inv_pref_var, width=15, font=(FONT_FAMILY, 10)).pack(anchor="w")
        tk.Label(frame, text="Default Low-Stock Threshold", font=(FONT_FAMILY, 10)).pack(anchor="w", pady=(12, 2))
        self._low_stock_var = tk.StringVar(value=str(get_setting("business.low_stock_threshold_default", 5)))
        tk.Entry(frame, textvariable=self._low_stock_var, width=10, font=(FONT_FAMILY, 10)).pack(anchor="w")
        tk.Label(frame, text="Expiry Alert Window (days)", font=(FONT_FAMILY, 10)).pack(anchor="w", pady=(12, 2))
        self._expiry_days_var = tk.StringVar(value=str(get_setting("business.expiry_alert_days", 30)))
        tk.Entry(frame, textvariable=self._expiry_days_var, width=10, font=(FONT_FAMILY, 10)).pack(anchor="w")
        tk.Label(frame, text="Items expiring within this many days show up on the dashboard card.\n"
                             "(Used when the Expiry Tracking feature is enabled.)",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(anchor="w", pady=(2, 0))
        tk.Label(frame, text="License Sales Email (for upgrade requests)", font=(FONT_FAMILY, 10)).pack(anchor="w", pady=(12, 2))
        self._sales_email_var = tk.StringVar(value=get_setting("business.sales_email", "") or "")
        tk.Entry(frame, textvariable=self._sales_email_var, width=32, font=(FONT_FAMILY, 10)).pack(anchor="w")
        tk.Label(frame, text="Where 'Request Upgrade / Extend' in the License tab sends its email.",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(anchor="w", pady=(2, 0))
        self._show_tax_recpt_var = tk.BooleanVar(value=get_setting("business.show_tax_breakdown_on_receipt", True))
        tk.Checkbutton(frame, text="Show tax breakdown on receipt", variable=self._show_tax_recpt_var, bg=CARD_BG, font=(FONT_FAMILY, 10), activebackground=CARD_BG).pack(anchor="w", pady=(14, 0))
        tk.Button(frame, text="Save Business Settings", command=self._save_business, font=(FONT_FAMILY, 10, "bold")).pack(anchor="w", pady=(16, 0))

    def _save_business(self):
        set_setting("business.invoice_no_prefix", self._inv_pref_var.get().strip())
        try:
            set_setting("business.low_stock_threshold_default", int(self._low_stock_var.get()))
        except ValueError:
            pass
        try:
            set_setting("business.expiry_alert_days", int(self._expiry_days_var.get()))
        except ValueError:
            pass
        set_setting("business.sales_email", self._sales_email_var.get().strip())
        set_setting("business.show_tax_breakdown_on_receipt", bool(self._show_tax_recpt_var.get()))
        messagebox.showinfo("Saved", "Business settings saved")

    def _vertical_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Vertical  ")

        # Current vertical info card
        from utils.verticals import VERTICALS, ALL_FEATURES
        cur_key = get_setting("vertical", "general")
        cur = VERTICALS.get(cur_key, VERTICALS["general"])
        info = tk.Frame(frame, bg=CARD_BG, highlightbackground="#CBD5E1",
                        highlightthickness=1, padx=14, pady=12)
        info.pack(fill=tk.X, pady=(0, 14))
        top = tk.Frame(info, bg=CARD_BG); top.pack(fill=tk.X)
        tk.Label(top, text=cur.get("icon", ""), font=("Segoe UI Emoji", 20),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(side=tk.LEFT, padx=(0, 10))
        tk.Label(top, text=f"{cur.get('name', cur_key)}  ·  active",
                 font=(FONT_FAMILY, FONT_SIZE_LG, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(side=tk.LEFT)
        tk.Label(info, text=cur.get("description", ""),
                 font=(FONT_FAMILY, 9), bg=CARD_BG, fg=TEXT_SECONDARY,
                 wraplength=620, justify="left").pack(anchor="w", pady=(6, 4))
        feat_names = {fid: name for fid, name, _ in ALL_FEATURES}
        primary = [feat_names.get(f, f) for f in cur.get("primary_features", [])]
        if primary:
            tk.Label(info, text="Active features: " + ", ".join(primary),
                     font=(FONT_FAMILY, 9), bg=CARD_BG, fg=TEXT_MUTED,
                     wraplength=620, justify="left").pack(anchor="w")
        extra = cur.get("extra_stock_fields", []) + cur.get("extra_customer_fields", [])
        if extra:
            tk.Label(info, text="Extra fields: " + ", ".join(extra),
                     font=(FONT_FAMILY, 9), bg=CARD_BG, fg=TEXT_MUTED,
                     wraplength=620, justify="left").pack(anchor="w")

        tk.Label(frame, text="Business Type", font=(FONT_FAMILY, 12, "bold"), bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w", pady=(0, 8))
        tk.Label(frame, text="Choose your industry. This enables appropriate features and fields.", font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(anchor="w", pady=(0, 8))
        from utils.verticals import list_verticals
        current = get_setting("vertical", "general")
        self._v_sel_var = tk.StringVar(value=current)
        for key, name, icon, desc in list_verticals():
            tk.Radiobutton(frame, text=icon + "  " + name, variable=self._v_sel_var, value=key, bg=CARD_BG, font=(FONT_FAMILY, 10), selectcolor=CARD_BG, activebackground=CARD_BG).pack(anchor="w", pady=2)
        tk.Button(frame, text="Apply Vertical", command=self._apply_vertical, font=(FONT_FAMILY, 10, "bold")).pack(anchor="w", pady=(12, 0))
        tk.Label(frame, text="You can also re-run the setup wizard", font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(anchor="w", pady=(4, 0))
        tk.Button(frame, text="Re-Run Setup Wizard", command=self._run_wizard, font=(FONT_FAMILY, 10)).pack(anchor="w", pady=(2, 0))

    def _apply_vertical(self):
        vertical = self._v_sel_var.get()
        from utils.settings_helper import update_features_from_vertical
        update_features_from_vertical(vertical)
        self._app._apply_sidebar_visibility()
        # Rebuild EVERY page — new fields/cards/panels only appear on a fresh
        # page instance (Stock form extras, dashboard cards, Sales KOT panel…).
        self._app.reload_all_pages()
        messagebox.showinfo("Vertical Updated",
                            "Business type and features applied. Open Stock/Sales/Customers "
                            "to see the new fields and panels.")

    def _run_wizard(self):
        from ui.wizard import FirstRunWizard
        FirstRunWizard().run(self._app)
        self._app._apply_sidebar_visibility()
        self._app.reload_current_page()

    def _features_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Features  ")
        tk.Label(frame, text="Feature Flags", font=(FONT_FAMILY, 12, "bold"), bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w", pady=(0, 4))
        tk.Label(frame, text="Enable/disable individual features. Some features depend on the vertical.", font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(anchor="w", pady=(0, 8))
        from utils.verticals import ALL_FEATURES
        self._feat_vars = {}
        for fid, name, desc in ALL_FEATURES:
            row = tk.Frame(frame, bg=CARD_BG)
            row.pack(fill="x", pady=2)
            var = tk.BooleanVar(value=get_setting("feature_flags." + fid, False))
            self._feat_vars[fid] = var
            tk.Checkbutton(row, text=name, variable=var, bg=CARD_BG, font=(FONT_FAMILY, 10, "bold"), activebackground=CARD_BG, selectcolor=CARD_BG).pack(side=tk.LEFT)
            tk.Label(row, text=desc, font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(side=tk.LEFT, padx=(8, 0))
        tk.Button(frame, text="Save Feature Settings", command=self._save_features, font=(FONT_FAMILY, 10, "bold")).pack(anchor="w", pady=(12, 0))

    def _save_features(self):
        for fid, var in self._feat_vars.items():
            set_setting("feature_flags." + fid, bool(var.get()))
        self._app._apply_sidebar_visibility()
        self._app.reload_all_pages()
        messagebox.showinfo("Saved", "Feature flags saved and applied to all pages.")

    def _tax_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Tax  ")
        row_frame = tk.Frame(frame, bg=CARD_BG)
        row_frame.pack(fill="x", pady=(0, 8))
        tk.Label(row_frame, text="Country / Tax System", font=(FONT_FAMILY, 12, "bold"), bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        country = get_setting("country", "India")
        self._c_var = tk.StringVar(value=country)
        for name in ["India", "Nepal", "None"]:
            tk.Radiobutton(row_frame, text=name, variable=self._c_var, value=name, bg=CARD_BG, font=(FONT_FAMILY, 10), selectcolor=CARD_BG, activebackground=CARD_BG, command=self._on_country_change).pack(anchor="w", pady=2)
        self._tax_enabled_var = tk.BooleanVar(value=get_setting("feature_flags.tax_system", True))
        tk.Checkbutton(frame, text="Enable Tax (show on invoices)", variable=self._tax_enabled_var, bg=CARD_BG, font=(FONT_FAMILY, 10), activebackground=CARD_BG, selectcolor=CARD_BG).pack(anchor="w", pady=(4, 8))
        self._tax_rate_frame = tk.Frame(frame, bg=CARD_BG)
        self._tax_rate_frame.pack(fill="x")
        tk.Label(self._tax_rate_frame, text="Default Tax Rate (%)", font=(FONT_FAMILY, 10, "bold"), bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        self._rate_var = tk.StringVar(value=str(get_setting("tax.default_rate_percent", 18)))
        tk.Entry(self._tax_rate_frame, textvariable=self._rate_var, width=8, font=(FONT_FAMILY, 10)).pack(anchor="w", pady=(2, 0))
        tk.Button(frame, text="Save Tax Settings", command=self._save_tax, font=(FONT_FAMILY, 10, "bold")).pack(anchor="w", pady=(12, 0))

    def _on_country_change(self):
        val = self._c_var.get()
        if val == "India":
            self._tax_enabled_var.set(True)
            self._rate_var.set("18")
        elif val == "Nepal":
            self._tax_enabled_var.set(True)
            self._rate_var.set("13")
        elif val == "None":
            self._tax_enabled_var.set(False)
            self._rate_var.set("0")
        if val == "India":
            set_setting("currency_symbol", "\u20B9")
        elif val == "Nepal":
            set_setting("currency_symbol", "Rs.")
        else:
            set_setting("currency_symbol", "$")

    def _save_tax(self):
        set_setting("country", self._c_var.get())
        set_setting("feature_flags.tax_system", bool(self._tax_enabled_var.get()))
        try:
            set_setting("tax.default_rate_percent", float(self._rate_var.get()))
        except ValueError:
            pass
        set_setting("first_run", False)
        self._app.reload_current_page()
        messagebox.showinfo("Saved", "Tax settings saved")

    # ========== NOTIFICATIONS TAB ==========
    def _notifications_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Notifications  ")

        from utils import notifier
        from utils.license import license_mgr

        canvas = tk.Canvas(frame, bg=CARD_BG, highlightthickness=0)
        vsb = ttk.Scrollbar(frame, orient="vertical", command=canvas.yview)
        body = tk.Frame(canvas, bg=CARD_BG)
        body.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=body, anchor="nw")
        canvas.configure(yscrollcommand=vsb.set)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)

        tk.Label(body, text="Notifications", font=(FONT_FAMILY, 12, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        tk.Label(body, text="Email customers their receipts and payment reminders. "
                            "Customers opt in per record (Customers page).",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG,
                 wraplength=640, justify="left").pack(anchor="w", pady=(0, 10))

        if not license_mgr.has_feature("sale_notifications"):
            tk.Label(body, text=("⚠  Email notifications are available in Basic/Pro/Enterprise plans. "
                                 "Activate a license in the License tab to enable sending."),
                     font=(FONT_FAMILY, 9, "bold"), fg=WARNING_COLOR, bg=CARD_BG,
                     wraplength=640, justify="left").pack(anchor="w", pady=(0, 10))

        self._notif_enabled_var = tk.BooleanVar(value=get_setting("notifications.enabled", True))
        tk.Checkbutton(body, text="Enable notifications (master switch)",
                       variable=self._notif_enabled_var, bg=CARD_BG,
                       font=(FONT_FAMILY, 10, "bold"), activebackground=CARD_BG,
                       selectcolor=CARD_BG).pack(anchor="w")

        tk.Label(body, text="Events", font=(FONT_FAMILY, 11, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w", pady=(14, 2))
        self._notif_sale_email_var = tk.BooleanVar(value=get_setting("notifications.on_sale_email", True))
        tk.Checkbutton(body, text="Email receipt after a sale (when the customer opted in)",
                       variable=self._notif_sale_email_var, bg=CARD_BG,
                       font=(FONT_FAMILY, 10), activebackground=CARD_BG,
                       selectcolor=CARD_BG).pack(anchor="w")
        self._notif_sale_wa_var = tk.BooleanVar(value=get_setting("notifications.on_sale_whatsapp", False))
        tk.Checkbutton(body, text="Open WhatsApp with receipt after a sale (you press Send, no API key needed)",
                       variable=self._notif_sale_wa_var, bg=CARD_BG,
                       font=(FONT_FAMILY, 10), activebackground=CARD_BG,
                       selectcolor=CARD_BG).pack(anchor="w")
        self._notif_reminders_var = tk.BooleanVar(value=get_setting("notifications.enable_reminders", True))
        tk.Checkbutton(body, text="Allow payment reminders (from Dashboard and Sales page)",
                       variable=self._notif_reminders_var, bg=CARD_BG,
                       font=(FONT_FAMILY, 10), activebackground=CARD_BG,
                       selectcolor=CARD_BG).pack(anchor="w")

        ttk.Separator(body, orient="horizontal").pack(fill=tk.X, pady=14)

        tk.Label(body, text="Email (SMTP)", font=(FONT_FAMILY, 11, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        tk.Label(body, text="Works with Gmail (App Password), Outlook, Zoho, or any SMTP server. "
                            "The password is stored encrypted (Windows DPAPI) — never in plain text.",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG,
                 wraplength=640, justify="left").pack(anchor="w", pady=(0, 8))

        grid = tk.Frame(body, bg=CARD_BG)
        grid.pack(fill=tk.X)

        self._smtp_host_var = tk.StringVar(value=get_setting("notifications.smtp_host", "") or "")
        self._smtp_port_var = tk.StringVar(value=str(get_setting("notifications.smtp_port", 587)))
        self._smtp_user_var = tk.StringVar(value=get_setting("notifications.smtp_user", "") or "")
        self._smtp_pass_var = tk.StringVar(value=notifier.load_smtp_password())
        self._smtp_from_name_var = tk.StringVar(value=get_setting("notifications.smtp_from_name", APP_NAME) or APP_NAME)
        self._smtp_from_email_var = tk.StringVar(value=get_setting("notifications.smtp_from_email", "") or "")
        self._smtp_tls_var = tk.BooleanVar(value=get_setting("notifications.smtp_use_tls", True))

        rows = [
            ("SMTP Host", self._smtp_host_var, 28, "e.g. smtp.gmail.com"),
            ("SMTP Port", self._smtp_port_var, 8, "587 (TLS) or 465 (SSL)"),
            ("Username", self._smtp_user_var, 28, "usually your email address"),
            ("Password", self._smtp_pass_var, 28, "stored encrypted"),
            ("From Name", self._smtp_from_name_var, 28, "shown as the sender name"),
            ("From Email", self._smtp_from_email_var, 28, "receipts are sent from this address"),
        ]
        for i, (label, var, width, hint) in enumerate(rows):
            tk.Label(grid, text=label, font=(FONT_FAMILY, 10),
                     bg=CARD_BG, fg=TEXT_PRIMARY).grid(row=i, column=0, sticky="w", pady=3, padx=(0, 10))
            show = "●" if label == "Password" else None
            tk.Entry(grid, textvariable=var, width=width, font=(FONT_FAMILY, 10),
                     show=show).grid(row=i, column=1, sticky="w", pady=3)
            tk.Label(grid, text=hint, font=(FONT_FAMILY, 8), fg=TEXT_MUTED,
                     bg=CARD_BG).grid(row=i, column=2, sticky="w", padx=(8, 0))

        tk.Checkbutton(body, text="Use TLS (STARTTLS) — recommended",
                       variable=self._smtp_tls_var, bg=CARD_BG,
                       font=(FONT_FAMILY, 10), activebackground=CARD_BG,
                       selectcolor=CARD_BG).pack(anchor="w", pady=(8, 0))

        status_text = ("✓  Email is configured" if notifier.is_email_configured()
                       else "✗  Email is not configured yet (fill host, username and From Email)")
        self._notif_status_lbl = tk.Label(body, text=status_text,
                                          font=(FONT_FAMILY, 9, "bold"),
                                          fg=(SUCCESS_COLOR if notifier.is_email_configured() else TEXT_MUTED),
                                          bg=CARD_BG)
        self._notif_status_lbl.pack(anchor="w", pady=(10, 0))

        btn_row = tk.Frame(body, bg=CARD_BG)
        btn_row.pack(anchor="w", pady=(10, 0))
        tk.Button(btn_row, text="Save Notification Settings", font=(FONT_FAMILY, 10, "bold"),
                  command=self._save_notifications).pack(side=tk.LEFT, padx=(0, 8))
        tk.Button(btn_row, text="Test Connection", font=(FONT_FAMILY, 10),
                  command=self._test_smtp).pack(side=tk.LEFT, padx=(0, 8))
        tk.Button(btn_row, text="Send Test Email to Myself", font=(FONT_FAMILY, 10),
                  command=self._send_test_email).pack(side=tk.LEFT)

        ttk.Separator(body, orient="horizontal").pack(fill=tk.X, pady=14)
        tk.Label(body, text="WhatsApp", font=(FONT_FAMILY, 11, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        tk.Label(body, text="WhatsApp needs no configuration: the app opens wa.me with the receipt "
                            "pre-filled and you press Send in the browser tab.",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG,
                 wraplength=640, justify="left").pack(anchor="w", pady=(0, 6))

        ttk.Separator(body, orient="horizontal").pack(fill=tk.X, pady=14)
        from utils import portal
        tk.Label(body, text="Customer Portal (online)", font=(FONT_FAMILY, 11, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        tk.Label(body, text="Customers sign in with their phone + PIN to see their live account "
                            "statement. Give access from Customers → 'Portal Access'.",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG,
                 wraplength=640, justify="left").pack(anchor="w", pady=(0, 6))
        self._portal_enabled_var = tk.BooleanVar(value=get_setting("portal.enabled", True))
        tk.Checkbutton(body, text="Enable customer portal sync", variable=self._portal_enabled_var,
                       bg=CARD_BG, font=(FONT_FAMILY, 10), activebackground=CARD_BG,
                       selectcolor=CARD_BG).pack(anchor="w")
        link_row = tk.Frame(body, bg=CARD_BG)
        link_row.pack(fill=tk.X, pady=(6, 0))
        tk.Label(link_row, text=f"Portal URL: {portal.portal_url()}",
                 font=("Consolas", 9), bg=CARD_BG, fg=TEXT_SECONDARY).pack(side=tk.LEFT)
        tk.Button(link_row, text="Copy Portal Link", font=(FONT_FAMILY, 9),
                  command=self._copy_portal_link).pack(side=tk.LEFT, padx=(10, 0))
        tk.Button(link_row, text="Sync All Portal Customers Now", font=(FONT_FAMILY, 9),
                  command=self._sync_portal_now).pack(side=tk.LEFT, padx=(8, 0))
        self._portal_status_lbl = tk.Label(body, text="", font=(FONT_FAMILY, 9),
                                           bg=CARD_BG, fg=TEXT_MUTED)
        self._portal_status_lbl.pack(anchor="w", pady=(4, 0))

    def _copy_portal_link(self):
        from utils import portal
        self._app.clipboard_clear()
        self._app.clipboard_append(portal.portal_url())
        try:
            self._app.toast.show("Portal link copied", "success", 2500)
        except Exception:
            pass

    def _sync_portal_now(self):
        import threading
        from utils import portal
        from utils.license import license_mgr
        if not license_mgr.has_feature("customer_notifications"):
            from utils.feature_gate import require_feature
            require_feature("customer_notifications", parent=self,
                            friendly_name="Customer Portal")
            return
        self._portal_status_lbl.config(text="Syncing…", fg=TEXT_MUTED)

        def _run():
            synced, failed = portal.sync_enabled_customers()
            msg = f"Synced {synced} customer(s)" + (f", {failed} failed" if failed else "")
            self.after(0, lambda: self._portal_status_lbl.config(
                text=msg, fg=(SUCCESS_COLOR if failed == 0 else DANGER_COLOR)))
            self.after(0, lambda: self._app.toast.show(
                f"Portal sync: {msg}", "success" if failed == 0 else "warning", 4000))

        self._save_notifications()
        threading.Thread(target=_run, daemon=True).start()

    def _save_notifications(self):
        from utils import notifier
        set_setting("notifications.enabled", bool(self._notif_enabled_var.get()))
        set_setting("notifications.on_sale_email", bool(self._notif_sale_email_var.get()))
        set_setting("notifications.on_sale_whatsapp", bool(self._notif_sale_wa_var.get()))
        set_setting("notifications.enable_reminders", bool(self._notif_reminders_var.get()))
        set_setting("portal.enabled", bool(self._portal_enabled_var.get()))
        set_setting("notifications.smtp_host", self._smtp_host_var.get().strip())
        try:
            set_setting("notifications.smtp_port", int(self._smtp_port_var.get().strip() or 587))
        except ValueError:
            set_setting("notifications.smtp_port", 587)
        set_setting("notifications.smtp_user", self._smtp_user_var.get().strip())
        set_setting("notifications.smtp_from_name", self._smtp_from_name_var.get().strip())
        set_setting("notifications.smtp_from_email", self._smtp_from_email_var.get().strip())
        set_setting("notifications.smtp_use_tls", bool(self._smtp_tls_var.get()))
        notifier.save_smtp_password(self._smtp_pass_var.get())
        # Remove any legacy plaintext password from settings.json.
        if get_setting("notifications.smtp_password", None) is not None:
            set_setting("notifications.smtp_password", "")
        configured = notifier.is_email_configured()
        self._notif_status_lbl.config(
            text=("✓  Email is configured" if configured else "✗  Email is not configured yet"),
            fg=(SUCCESS_COLOR if configured else TEXT_MUTED))
        messagebox.showinfo("Saved", "Notification settings saved")

    def _test_smtp(self):
        import threading
        from utils import notifier
        self._notif_status_lbl.config(text="Testing connection…", fg=TEXT_MUTED)

        def _run():
            ok, msg = notifier.test_smtp_connection()
            color = SUCCESS_COLOR if ok else DANGER_COLOR
            def _done():
                self._notif_status_lbl.config(text=("✓  " if ok else "✗  ") + msg, fg=color)
                try:
                    self._app.toast.show(msg, "success" if ok else "error", 5000)
                except Exception:
                    pass
            self.after(0, _done)

        threading.Thread(target=_run, daemon=True).start()

    def _send_test_email(self):
        import threading
        from utils import notifier
        to = self._smtp_from_email_var.get().strip()
        if not to:
            messagebox.showwarning("Test Email", "Fill in 'From Email' first — the test email goes there.")
            return
        ok_cfg, cfg_msg = True, ""
        if not notifier.is_email_configured():
            messagebox.showwarning("Test Email", "SMTP is not fully configured yet. Save settings first.")
            return

        def _run():
            ok, msg = notifier.send_email(
                to_email=to,
                subject="Accounting Pro — test email",
                body_text="This is a test email from Accounting Pro.\nIf you can read this, SMTP is working.",
            )
            def _done():
                if ok:
                    try:
                        self._app.toast.show(msg, "success", 5000)
                    except Exception:
                        pass
                messagebox.showinfo("Test Email" if ok else "Test Failed", msg)
            self.after(0, _done)

        threading.Thread(target=_run, daemon=True).start()

    # ========== USERS TAB ==========
    def _users_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Users  ")

        from utils.auth import auth_manager, ROLES
        tk.Label(frame, text="User Management", font=(FONT_FAMILY, 12, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        tk.Label(frame, text="Manage who can access this system.",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(anchor="w", pady=(0, 10))

        info_frame = tk.Frame(frame, bg=CARD_BG)
        info_frame.pack(fill=tk.X, pady=(0, 10))
        current_user = auth_manager.get_current_user() or "Not logged in"
        current_role = auth_manager.get_current_role() or "viewer"
        tk.Label(info_frame, text=f"Current User: {current_user}  |  Role: {current_role.capitalize()}",
                 font=(FONT_FAMILY, 10), bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")

        # User list
        list_frame = tk.Frame(frame, bg=CARD_BG)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=4)

        tree = ttk.Treeview(list_frame, columns=["user", "role", "name", "created"],
                           show="headings", height=6)
        tree.heading("user", text="Username")
        tree.heading("role", text="Role")
        tree.heading("name", text="Display Name")
        tree.heading("created", text="Created At")
        tree.column("user", width=100)
        tree.column("role", width=80)
        tree.column("name", width=120)
        tree.column("created", width=140)
        for u in auth_manager.list_users():
            tree.insert("", tk.END, values=[
                u.get("username", "") or "",
                (u.get("role", "") or "viewer").capitalize(),
                u.get("display_name", "") or "",
                u.get("created_at", "") or "",
            ])
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll = ttk.Scrollbar(list_frame, orient="vertical", command=tree.yview)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)
        tree.configure(yscrollcommand=scroll.set)

        btn_frame = tk.Frame(frame, bg=CARD_BG)
        btn_frame.pack(fill=tk.X, pady=8)

        def open_user_management():
            from ui.login_dialog import UserManagementDialog
            UserManagementDialog(self._app).show()
            # Reload the tab
            for i in range(self._notebook.index("end")):
                if self._notebook.tab(i, "text") == "  Users  ":
                    self._notebook.forget(i)
                    break
            self._users_tab(self._notebook)

        ttk.Button(btn_frame, text="Manage Users (Add / Delete / Change Password)",
                  command=open_user_management).pack(side=tk.LEFT, padx=2)

    # ========== AUDIT LOG TAB ==========
    def _audit_log_tab(self, notebook):
        frame = ttk.Frame(notebook, padding=20)
        notebook.add(frame, text="  Audit Log  ")

        tk.Label(frame, text="Audit Trail", font=(FONT_FAMILY, 12, "bold"),
                 bg=CARD_BG, fg=TEXT_PRIMARY).pack(anchor="w")
        tk.Label(frame, text="All CREATE / UPDATE / DELETE actions are logged with user and timestamp.",
                 font=(FONT_FAMILY, 9), fg=TEXT_MUTED, bg=CARD_BG).pack(anchor="w", pady=(0, 10))

        list_frame = tk.Frame(frame, bg=CARD_BG)
        list_frame.pack(fill=tk.BOTH, expand=True)

        tree = ttk.Treeview(list_frame, columns=["time", "user", "action", "entity", "id"],
                           show="headings", height=12)
        tree.heading("time", text="Timestamp")
        tree.heading("user", text="User")
        tree.heading("action", text="Action")
        tree.heading("entity", text="Entity")
        tree.heading("id", text="Record ID")
        tree.column("time", width=150)
        tree.column("user", width=80)
        tree.column("action", width=70)
        tree.column("entity", width=80)
        tree.column("id", width=60)

        def load_audit():
            for item in tree.get_children():
                tree.delete(item)
            try:
                from database import excel_db as db
                if db._active_file:
                    from openpyxl import load_workbook
                    wb = load_workbook(db._active_file)
                    if "AuditLog" in wb.sheetnames:
                        ws = wb["AuditLog"]
                        for row in ws.iter_rows(min_row=2, values_only=True):
                            if any(v for v in row):
                                tree.insert("", 0, values=list(row)[:5])
                    wb.close()
            except Exception:
                pass

        load_audit()
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll = ttk.Scrollbar(list_frame, orient="vertical", command=tree.yview)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)
        tree.configure(yscrollcommand=scroll.set)

        ttk.Button(frame, text="Refresh", command=load_audit).pack(anchor="w", pady=6)